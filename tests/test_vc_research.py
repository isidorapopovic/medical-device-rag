"""Offline synthetic research fixtures; no real performance/clinical annotations."""
import io
import json
import socket

import pymupdf
import pytest
import requests
from streamlit.testing.v1 import AppTest

from vc_engine.__main__ import main
from vc_engine.comparables import company_candidates
from vc_engine.evaluation import DEFAULT_GOLD, evaluate_file, evaluate_labels
from vc_engine.fda_client import FDAPage
from vc_engine.ingest import ingest
from vc_engine.memo import SECTIONS, audit_claims, build_research, export_research, render_markdown
from vc_engine.regulatory_search import normalize
from vc_engine.regulatory_store import FDAStore
from vc_engine.scoring import DEFAULT_WEIGHTS, scorecard
from vc_engine.startup_profile import profile_from_dict, profile_payload
from vc_engine.startup_uploads import ingest_startup_upload
from test_app_services import workspace
from test_vc_engine import write_csv


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("VC research must work without network or Ollama")
    monkeypatch.setattr(socket, "create_connection", forbidden)
    monkeypatch.setattr(requests.Session, "get", forbidden)
    monkeypatch.setattr(requests.Session, "post", forbidden)


@pytest.fixture
def corpus(tmp_path):
    paths = {name: tmp_path / (name + ".csv") for name in ("companies", "investments", "documents")}
    write_csv(paths["companies"], "companies", [dict(company="Synthetic Heart", clinical_use_case="Cardiac ultrasound triage",
        selected_device_or_software="Cardiac ultrasound", headquarters_country="Germany", clinical_category="Cardiology",
        dealroom_reported_total_funding_usd_millions="12.5", funding_source_url="https://example.invalid/funding", research_date="2026-10-17")])
    write_csv(paths["investments"], "investments", [dict(investor="Synthetic Fund", company="Synthetic Heart", round_stage="Seed",
        reported_round_amount="10", reported_round_currency="EUR million", amount_type="Total financing round",
        round_announcement_source="https://example.invalid/round", as_of_date="2026-10-17"),
        dict(investor="Synthetic Fund", company="Investment Only Placeholder")])
    write_csv(paths["documents"], "source_documents", [dict(source_id="SYN-ONLY", fund_name="Synthetic Fund",
        document_title="Synthetic local catalog", direct_document_or_page_url="https://example.invalid/catalog")])
    db, fda_db = tmp_path / "vc.sqlite", tmp_path / "fda.sqlite"
    ingest(**paths, db=db, report_path=tmp_path / "quality.json")
    raws = [dict(k_number="K000001", device_name="Synthetic Cardiac ultrasound triage", applicant="Synthetic Applicant",
        decision_code="SESE", decision_description="SYNTHETIC Substantially Equivalent"),
        dict(k_number="K000002", device_name="Synthetic Cardiac ultrasound triage", decision_code="NSE",
             decision_description="SYNTHETIC Not Substantially Equivalent")]
    page = FDAPage(dict(results=raws), "https://api.fda.gov/device/510k.json?SYNTHETIC_ONLY", "2026-10-20T00:00:00+00:00", True)
    with FDAStore(fda_db) as store:
        store.save({}, page.payload, page.source_url, page.fetched_at, [normalize(raw, page) for raw in raws])
    return dict(db=db, fda_db=fda_db)


def profile(**updates):
    payload = dict(name="Synthetic Startup", intended_use="Cardiac ultrasound triage", input_modality="Ultrasound",
        clinical_indication="Cardiac", patient_population="Adults")
    payload.update(updates)
    return profile_from_dict(payload)


def test_offline_three_csv_to_memo_and_unique_exports(corpus, tmp_path):
    bundle = build_research(profile(), **corpus)
    assert bundle["memo"]["generation_mode"] == "deterministic_offline"
    assert bundle["memo"]["citation_audit"]["claims_withheld"] == 0
    assert bundle["memo"]["citation_audit"]["citation_validity"] == 1
    assert {row["kind"] for row in bundle["comparables"]["results"]} == {"company", "fda_510k_record"}
    text = render_markdown(bundle["memo"])
    assert all("## " + section in text for section in SECTIONS)
    assert "SYNTHETIC Not Substantially Equivalent" in text
    assert "not an investment recommendation" in text
    first = export_research(bundle, root=tmp_path / "exports")
    second = export_research(bundle, root=tmp_path / "exports")
    assert first != second
    assert {p.name for p in first.iterdir()} >= {"memo.md", "memo.json", "comparables.json", "scorecard.json"}
    saved = json.loads((first / "memo.json").read_text(encoding="utf-8"))
    assert all(ref["source_id"] in saved["sources"] for claim in saved["claims"] for ref in claim["citations"])
    assert saved["profile_review_status"] == "pending_review"
    assert all(row["score"] is None for row in bundle["scorecard"]["criteria"])


def test_funding_keeps_amount_types_currencies_and_sources(corpus):
    bundle = build_research(profile(), **corpus)
    row, = bundle["memo"]["funding_table"]
    assert (row["round_amount"]["amount"], row["round_amount"]["currency"]) == (10000000, "EUR")
    assert (row["cumulative_funding"]["amount"], row["cumulative_funding"]["currency"]) == (12500000, "USD")
    assert row["individual_check"]["amount"] is None
    for key in ("round_amount", "cumulative_funding", "individual_check"):
        assert {"amount_type", "currency", "source_url", "as_of_date"} <= row[key].keys()
    assert row["round_amount"]["source_url"] == "https://example.invalid/round"
    assert row["cumulative_funding"]["source_url"] == "https://example.invalid/funding"
    assert "type: Total financing round" in render_markdown(bundle["memo"])


@pytest.mark.parametrize("reference", [dict(source_id="invented", field="value", quote="literal"),
    dict(source_id="known", field="value", quote="fabricated"), dict(source_id="known", field="wrong", quote="literal"),
    dict(source_id=["invalid"], field="value", quote="literal"), "bad-object"])
def test_bad_citations_withhold_factual_claim(reference):
    claims, result = audit_claims([dict(section=SECTIONS[0], text="Unsupported investment fact", kind="public_source_fact",
        status="cached", citations=[reference])], {"known": {"fields": {"value": "literal"}}})
    assert result["claims_withheld"] == 1
    assert result["citation_validity"] == 0
    assert claims[0]["status"] == "needs_review" and claims[0]["citations"] == []
    assert "Unsupported" not in claims[0]["text"]


def test_citationless_fact_withheld_and_unknown_gap_retained():
    rows, audit = audit_claims([dict(kind="public_source_fact", text="Revenue invented", status="cached"),
                              dict(kind="analyst_assessment", text="Unknown revenue", status="unknown")], {})
    assert rows[0]["status"] == "needs_review" and rows[1]["text"] == "Unknown revenue"
    assert audit["citation_validity"] is None


def test_export_reaudits_mutated_claims_and_withholds_funding(corpus, tmp_path):
    bundle = build_research(profile(name="../../CON:<script>"), **corpus)
    row = next(claim for claim in bundle["memo"]["claims"] if claim.get("claim_id"))
    row["citations"].append(dict(source_id="invented", field="text", quote="false"))
    saved = export_research(bundle, root=tmp_path)
    assert saved.parent == tmp_path.resolve() and saved.name.startswith("profile-")
    memo = json.loads((saved / "memo.json").read_text(encoding="utf-8"))
    assert memo["funding_table"] == [] and memo["citation_audit"]["claims_withheld"] == 1
    assert "<script>" not in (saved / "memo.md").read_text(encoding="utf-8")
    assert bundle["memo"]["funding_table"]  # Export does not mutate its input.


def rating(score=4, **updates):
    row = dict(score=score, reason="Synthetic human assessment only", evidence_ids=["known"], confidence="low",
               missing_inputs=[], reviewed_by_human=True, reviewer="Synthetic Reviewer", evidence_verified=True)
    row.update(updates)
    return row


def test_partial_rubric_mean_coverage_and_guard():
    card = scorecard({"ratings": {"regulatory_feasibility": rating()}}, source_ids={"known"}, enabled=True)
    assert card["coverage"] == .2 and card["rated_mean_0_to_5"] == 4
    assert card["screening_score_0_to_5"] is None and not card["ranking_allowed"]
    assert all(row["score"] is None for row in card["criteria"] if row["category"] != "regulatory_feasibility")


@pytest.mark.parametrize("changes", [dict(evidence_verified=False), dict(reviewed_by_human=False), dict(reviewer=""),
    dict(evidence_ids=["invented"]), dict(evidence_ids=[]), dict(reason="")])
def test_unverified_criterion_never_becomes_numeric(changes):
    card = scorecard({"ratings": {"team_execution": rating(**changes)}}, source_ids={"known"}, enabled=True)
    assert card["coverage"] == 0 and card["rated_mean_0_to_5"] is None


def test_manual_full_rubric_accepts_zero_and_can_disable():
    config = {"minimum_coverage": 1, "ratings": {key: rating(score=0) for key in DEFAULT_WEIGHTS}}
    card = scorecard(config, source_ids={"known"}, enabled=True)
    assert card["coverage"] == 1 and card["screening_score_0_to_5"] == 0 and card["ranking_allowed"]
    assert scorecard(config, source_ids={"known"}, enabled=False)["coverage"] == 0


@pytest.mark.parametrize("config", [{"weights": {"team_execution": 1}}, {"minimum_coverage": .1},
    {"ratings": {"team_execution": {"score": 7}}}, {"ratings": {"team_execution": {"score": True}}},
    {"ratings": {"team_execution": {"confidence": []}}}])
def test_invalid_rubric_rejected(config):
    with pytest.raises(ValueError):
        scorecard(config, enabled=True)


def test_empty_gold_has_no_fabricated_metrics():
    result = evaluate_file(DEFAULT_GOLD)
    assert result["status"] == "awaiting_human_annotations"
    assert all(result[key] is None for key in ("recall_at_5", "precision_at_5", "extraction_field_accuracy", "citation_validity"))


@pytest.mark.parametrize("memo", [{}, [], {"claims": [], "sources": []}])
def test_invalid_memo_evaluation_rejected_cleanly(memo):
    with pytest.raises(ValueError):
        evaluate_labels(dict(schema_version=1, examples=[]), memo=memo)


def annotated(corpus):
    candidate, = [row for row in company_candidates(corpus["db"], 100) [0] if row["name"] == "Synthetic Heart"]
    return dict(schema_version=1, examples=[dict(id="SYNTHETIC-UNIT-ONLY", profile=profile_payload(profile()),
        candidate_kind="company", relevant_ids=[candidate["source_id"]], reviewed_by="Synthetic Reviewer",
        reviewed_on="2026-10-21", annotation_notes="Synthetic test, not real ground truth",
        expected_fields={"name": "Synthetic Startup", "input_modality": "CT", "technology": None})])


def test_manual_label_metrics_and_literal_citation_accuracy(corpus):
    bundle = build_research(profile(), **corpus)
    result = evaluate_labels(annotated(corpus), **corpus, memo=bundle["memo"])
    assert result["recall_at_5"] == 1 and result["precision_at_5"] == .2
    assert result["extraction_field_accuracy"] == pytest.approx(2 / 3)
    assert result["citation_validity"]["citation_validity"] == 1


@pytest.mark.parametrize("change", ["invented", "no_reviewer", "bad_date", "no_profile", "wrong_kind", "invalid_kind_type"])
def test_annotations_must_be_human_attributed_and_valid(corpus, change):
    payload = annotated(corpus)
    row = payload["examples"][0]
    if change == "invented":
        row["relevant_ids"] = ["K999999"]
    elif change == "no_reviewer":
        row["reviewed_by"] = ""
    elif change == "bad_date":
        row["reviewed_on"] = 15
    elif change == "no_profile":
        row.pop("profile")
    elif change == "invalid_kind_type":
        row["candidate_kind"] = []
    else:
        row["candidate_kind"] = "fda_510k_record"
    with pytest.raises(ValueError):
        evaluate_labels(payload, **corpus)


def test_local_pdf_extraction_retains_original_page_and_memo_citation(tmp_path):
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "Synthetic cover only")
    document.new_page().insert_text((72, 72), "Company name: Synthetic Deck\nIntended use: Cardiac ultrasound triage\nInput modality: Ultrasound")
    result = ingest_startup_upload("synthetic-deck.pdf", document.tobytes(), root=tmp_path / "uploads")
    document.close()
    parsed = result["profile"]
    assert parsed.name == "Synthetic Deck" and parsed.sources[1].page == 2
    assert parsed.review_status == "pending_review" and parsed.technology is None
    assert all(anchor.quote in next(source.text for source in parsed.sources if source.source_id == anchor.source_id) for anchor in parsed.field_evidence)
    bundle = build_research(parsed, db=tmp_path / "missing.sqlite", fda_db=tmp_path / "missing-fda.sqlite")
    text = render_markdown(bundle["memo"])
    assert "synthetic-deck.pdf" in text and "page 2" in text
    assert bundle["memo"]["citation_audit"]["claims_withheld"] == 0
    assert (result["directory"] / "original.pdf").exists()


def test_docx_conflicts_stay_unknown_and_duplicate_uploads_survive(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Company name: Synthetic Deck")
    doc.add_paragraph("Input modality: CT")
    doc.add_paragraph("Input modality: MRI")
    content = io.BytesIO()
    doc.save(content)
    first = ingest_startup_upload("synthetic.docx", content.getvalue(), root=tmp_path)
    second = ingest_startup_upload("synthetic.docx", content.getvalue(), root=tmp_path)
    assert first["directory"] != second["directory"]
    assert first["profile"].input_modality is None and any("Conflicting" in warning for warning in first["warnings"])


@pytest.mark.parametrize("filename,content", [("../deck.json", b"{}"), ("C:\\deck.pdf", b"pdf"), ("deck.exe", b"exe"),
    ("deck.pdf", b""), ("deck.json", b" " * (2 * 1024 * 1024 + 1)), ("deck.pdf", b" " * (20 * 1024 * 1024 + 1))],
    ids=["traversal", "absolute-path", "extension", "empty", "json-limit", "deck-limit"])
def test_upload_limits_and_path_traversal_rejected(tmp_path, filename, content):
    with pytest.raises(ValueError):
        ingest_startup_upload(filename, content, root=tmp_path)
    assert not list(tmp_path.iterdir())


def test_json_upload_review_invalidated_after_edit(tmp_path):
    content = json.dumps(profile_payload(profile_from_dict({"name": "Original"}, reviewer="Synthetic Reviewer"))).encode()
    result = ingest_startup_upload("profile.json", content, root=tmp_path)
    assert result["profile"].review_status == "user_reviewed"
    changed = profile_payload(result["profile"])
    changed["name"] = "Changed"
    assert profile_from_dict(changed).review_status == "pending_review"


def test_cli_memo_and_empty_evaluation_without_network(corpus, tmp_path, capsys):
    source = tmp_path / "startup.json"
    source.write_text(json.dumps(profile_payload(profile())), encoding="utf-8")
    options = ["--db", str(corpus["db"]), "--fda-db", str(corpus["fda_db"])]
    assert main(["memo", "--profile", str(source), "--output-root", str(tmp_path / "exports"), *options]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["coverage"] == 0
    assert main(["evaluate-vc", *options]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "awaiting_human_annotations"
    assert main(["evaluate-vc", "--labels", str(DEFAULT_GOLD), "--output", str(DEFAULT_GOLD), *options]) == 1


def dashboard(corpus):
    script = "from pathlib import Path\nfrom vc_engine.dashboard import render_vc_research\nrender_vc_research(db=Path(%r), fda_db=Path(%r))" % (str(corpus["db"]), str(corpus["fda_db"]))
    return AppTest.from_string(script, default_timeout=30).run()


def button(ui, label):
    return next(item for item in ui.button if item.label == label)


def test_streamlit_company_explorer_filters_and_omits_placeholder(corpus):
    ui = dashboard(corpus)
    assert not ui.exception and not ui.error
    assert any("1 company-catalog records" in caption.value for caption in ui.caption)
    ui.text_input(key="vc_company_query").set_value("not present").run()
    assert not ui.exception and not ui.error
    assert any("0 company-catalog records" in caption.value for caption in ui.caption)


def test_streamlit_evaluate_memo_score_guard_and_empty_gold(corpus):
    ui = dashboard(corpus)
    ui.radio(key="vc_view").set_value("Startup evaluator").run()
    ui.text_area(key="vc_editor").set_value(json.dumps(profile_payload(profile()))).run()
    button(ui, "Evaluate startup offline").click().run()
    assert not ui.exception and not ui.error
    assert ui.session_state["vc_bundle"]["memo"]["citation_audit"]["claims_withheld"] == 0
    ui.radio(key="vc_view").set_value("Research memo").run()
    ui.checkbox(key="vc_scoring_enabled").check().run()
    button(ui, "Generate research memo").click().run()
    assert not ui.exception and not ui.error
    assert next(item for item in ui.metric if item.label == "Human-rated rubric coverage").value == "0%"
    assert next(item for item in ui.metric if item.label == "Screening score (coverage guard)").value == "Withheld"
    ui.radio(key="vc_view").set_value("Evaluation").run()
    button(ui, "Evaluate human labels").click().run()
    assert not ui.exception and not ui.error
    assert ui.session_state["vc_evaluation"]["recall_at_5"] is None


def test_streamlit_bad_profile_fails_gracefully(corpus):
    ui = dashboard(corpus)
    ui.radio(key="vc_view").set_value("Startup evaluator").run()
    ui.text_area(key="vc_editor").set_value("invalid JSON").run()
    button(ui, "Evaluate startup offline").click().run()
    assert not ui.exception and ui.error


def test_streamlit_charts_show_supplied_human_label_metrics(corpus):
    ui = dashboard(corpus)
    ui.session_state["vc_evaluation"] = evaluate_labels(annotated(corpus), **corpus)
    ui.radio(key="vc_view").set_value("Evaluation").run()
    assert not ui.exception and not ui.error
    assert ui.session_state["vc_evaluation"]["precision_at_5"] == .2


def test_existing_entry_point_defaults_to_mdr_and_switches_to_vc(workspace, corpus, monkeypatch):
    from vc_engine import dashboard as screens
    render = screens.render_vc_research
    monkeypatch.setattr(screens, "render_vc_research", lambda: render(**corpus))
    ui = AppTest.from_string("import streamlit_app as app\napp.main()", default_timeout=30).run()
    assert not ui.exception and ui.selectbox(key="workspace").value == "MDR"
    assert any(item.label == "Index Documents" for item in ui.button)
    ui.selectbox(key="workspace").set_value("VC Research").run()
    assert not ui.exception and not ui.error
    assert any(item.value == "VC Research" for item in ui.title)


def test_startup_draft_survives_navigation_and_edits_clear_review(corpus):
    ui = dashboard(corpus)
    ui.radio(key="vc_view").set_value("Startup evaluator").run()
    text = json.dumps(profile_payload(profile(name="Synthetic retained draft")))
    ui.text_area(key="vc_editor").set_value(text).run()
    ui.text_input(key="vc_reviewer").set_value("Synthetic Reviewer").run()
    ui.checkbox(key="vc_review_confirmed").check().run()
    ui.text_area(key="vc_editor").set_value(text.replace("retained", "changed")).run()
    assert not ui.checkbox(key="vc_review_confirmed").value
    ui.radio(key="vc_view").set_value("Research memo").run()
    ui.radio(key="vc_view").set_value("Startup evaluator").run()
    assert not ui.exception and not ui.error
    assert "Synthetic changed draft" in ui.text_area(key="vc_editor").value
